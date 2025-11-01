"""
Template Engine - Generate legal documents from templates with merge fields
"""
from jinja2 import Template, Environment, FileSystemLoader
from docx import Document as DocxDocument
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict, List, Optional
from datetime import datetime
import os
from sqlalchemy.orm import Session
from models import Case, Document, Contact, Party
import json


class TemplateEngine:
    """Generate documents from templates with variable substitution"""

    def __init__(self, db: Session, case_id: int):
        self.db = db
        self.case_id = case_id
        self.case = db.query(Case).filter(Case.id == case_id).first()

        if not self.case:
            raise ValueError(f"Case {case_id} not found")

        # Setup template directory
        self.template_dir = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),
            'templates',
            'document_templates'
        )
        os.makedirs(self.template_dir, exist_ok=True)

        # Setup Jinja2 environment
        self.jinja_env = Environment(loader=FileSystemLoader(self.template_dir))

    def get_available_templates(self) -> List[Dict]:
        """Get list of available document templates"""
        templates = []

        for filename in os.listdir(self.template_dir):
            if filename.endswith('.json'):
                template_path = os.path.join(self.template_dir, filename)
                try:
                    with open(template_path, 'r') as f:
                        template_config = json.load(f)
                        templates.append({
                            'id': filename.replace('.json', ''),
                            'name': template_config.get('name'),
                            'description': template_config.get('description'),
                            'category': template_config.get('category'),
                            'fields': template_config.get('fields', [])
                        })
                except Exception as e:
                    print(f"Error loading template {filename}: {e}")

        return templates

    def generate_document_from_template(
        self,
        template_id: str,
        variables: Dict,
        output_format: str = 'docx'
    ) -> Dict:
        """
        Generate a document from a template

        Args:
            template_id: Template identifier
            variables: Dictionary of variable values
            output_format: 'docx' or 'txt'

        Returns:
            Dict with generated document info
        """
        # Load template configuration
        template_config_path = os.path.join(
            self.template_dir,
            f"{template_id}.json"
        )

        if not os.path.exists(template_config_path):
            return {'error': f'Template {template_id} not found'}

        with open(template_config_path, 'r') as f:
            template_config = json.load(f)

        # Merge case data with provided variables
        merged_variables = self._get_case_variables()
        merged_variables.update(variables)

        # Generate document based on format
        if output_format == 'docx':
            result = self._generate_docx(template_id, template_config, merged_variables)
        else:
            result = self._generate_text(template_id, template_config, merged_variables)

        if result.get('success'):
            # Save as document record
            document = Document(
                case_id=self.case_id,
                filename=result['filename'],
                file_path=result['file_path'],
                category=template_config.get('category', 'Other'),
                document_type='Generated Document',
                status='Draft',
                description=f"Generated from template: {template_config.get('name')}",
                file_size=os.path.getsize(result['file_path'])
            )

            self.db.add(document)
            self.db.commit()
            self.db.refresh(document)

            result['document_id'] = document.id

        return result

    def _generate_docx(
        self,
        template_id: str,
        template_config: Dict,
        variables: Dict
    ) -> Dict:
        """Generate Word document from template"""
        try:
            # Create document
            doc = DocxDocument()

            # Set up styles
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Arial'
            font.size = Pt(11)

            # Add title
            title = doc.add_heading(template_config.get('name', 'Document'), level=1)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            # Add metadata
            doc.add_paragraph(f"Case: {self.case.case_number}")
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y')}")
            doc.add_paragraph('')  # Blank line

            # Render template content
            template_text = template_config.get('content', '')
            jinja_template = Template(template_text)
            rendered_content = jinja_template.render(**variables)

            # Add rendered content (split by paragraphs)
            for paragraph in rendered_content.split('\n\n'):
                if paragraph.strip():
                    # Check for headings (lines starting with ##)
                    if paragraph.strip().startswith('##'):
                        heading_text = paragraph.strip().replace('##', '').strip()
                        doc.add_heading(heading_text, level=2)
                    else:
                        doc.add_paragraph(paragraph.strip())

            # Save document
            output_dir = os.path.join(
                os.path.dirname(os.path.dirname(__file__)),
                'documents',
                str(self.case_id),
                'generated'
            )
            os.makedirs(output_dir, exist_ok=True)

            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{template_id}_{timestamp}.docx"
            file_path = os.path.join(output_dir, filename)

            doc.save(file_path)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'template_id': template_id,
                'format': 'docx'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def _generate_text(
        self,
        template_id: str,
        template_config: Dict,
        variables: Dict
    ) -> Dict:
        """Generate text document from template"""
        try:
            # Render template content
            template_text = template_config.get('content', '')
            jinja_template = Template(template_text)
            rendered_content = jinja_template.render(**variables)

            # Add header
            header = f"""
{template_config.get('name', 'Document')}
{'=' * 60}

Case: {self.case.case_number}
Generated: {datetime.now().strftime('%d %B %Y')}

{'=' * 60}

"""
            full_content = header + rendered_content

            # Save document
            output_dir = os.path.join(
                os.path.dirname(os.path.dirname(__file__)),
                'documents',
                str(self.case_id),
                'generated'
            )
            os.makedirs(output_dir, exist_ok=True)

            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{template_id}_{timestamp}.txt"
            file_path = os.path.join(output_dir, filename)

            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(full_content)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'template_id': template_id,
                'format': 'txt'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def _get_case_variables(self) -> Dict:
        """Get standard variables from case data"""
        variables = {
            'case_number': self.case.case_number,
            'case_title': self.case.title,
            'case_type': self.case.case_type,
            'case_status': self.case.status,
            'mission_statement': self.case.mission_statement or '',
            'today': datetime.now().strftime('%d %B %Y'),
            'today_short': datetime.now().strftime('%d/%m/%Y')
        }

        # Add parties
        parties = self.db.query(Party).filter(Party.case_id == self.case_id).all()
        for party in parties:
            party_type = party.party_type.lower().replace(' ', '_')
            variables[f'{party_type}_name'] = party.name
            variables[f'{party_type}_org'] = party.organization or ''

        # Add contacts
        contacts = self.db.query(Contact).filter(Contact.case_id == self.case_id).all()
        if contacts:
            # Add first contact as primary
            primary_contact = contacts[0]
            variables['primary_contact_name'] = primary_contact.name
            variables['primary_contact_email'] = primary_contact.email or ''
            variables['primary_contact_phone'] = primary_contact.phone or ''

        return variables

    def create_template(
        self,
        name: str,
        description: str,
        category: str,
        content: str,
        fields: List[Dict]
    ) -> Dict:
        """
        Create a new document template

        Args:
            name: Template name
            description: Template description
            category: Document category
            content: Template content with Jinja2 variables
            fields: List of field definitions

        Returns:
            Dict with template info
        """
        # Generate template ID
        template_id = name.lower().replace(' ', '_').replace('-', '_')

        # Create template configuration
        template_config = {
            'name': name,
            'description': description,
            'category': category,
            'content': content,
            'fields': fields,
            'created': datetime.now().isoformat()
        }

        # Save template
        template_path = os.path.join(
            self.template_dir,
            f"{template_id}.json"
        )

        with open(template_path, 'w', encoding='utf-8') as f:
            json.dump(template_config, f, indent=2)

        return {
            'success': True,
            'template_id': template_id,
            'template_path': template_path
        }

    def get_template_preview(
        self,
        template_id: str,
        variables: Optional[Dict] = None
    ) -> Dict:
        """
        Get a preview of the rendered template

        Args:
            template_id: Template identifier
            variables: Optional variables (uses defaults if not provided)

        Returns:
            Dict with preview text
        """
        template_config_path = os.path.join(
            self.template_dir,
            f"{template_id}.json"
        )

        if not os.path.exists(template_config_path):
            return {'error': f'Template {template_id} not found'}

        with open(template_config_path, 'r') as f:
            template_config = json.load(f)

        # Merge variables
        merged_variables = self._get_case_variables()
        if variables:
            merged_variables.update(variables)

        # Render template
        template_text = template_config.get('content', '')
        jinja_template = Template(template_text)
        rendered_content = jinja_template.render(**merged_variables)

        return {
            'success': True,
            'preview': rendered_content,
            'template_name': template_config.get('name'),
            'fields_used': list(merged_variables.keys())
        }


def create_default_templates():
    """Create default legal document templates"""
    templates_dir = os.path.join(
        os.path.dirname(os.path.dirname(__file__)),
        'templates',
        'document_templates'
    )
    os.makedirs(templates_dir, exist_ok=True)

    # Witness Statement Template
    witness_statement = {
        'name': 'Witness Statement',
        'description': 'Standard witness statement for tribunal proceedings',
        'category': 'Evidence',
        'fields': [
            {'name': 'witness_name', 'label': 'Witness Name', 'type': 'text', 'required': True},
            {'name': 'witness_address', 'label': 'Witness Address', 'type': 'textarea', 'required': True},
            {'name': 'statement_date', 'label': 'Statement Date', 'type': 'date', 'required': True},
            {'name': 'statement_content', 'label': 'Statement Content', 'type': 'textarea', 'required': True}
        ],
        'content': """## WITNESS STATEMENT

**Name:** {{ witness_name }}

**Address:** {{ witness_address }}

**Case:** {{ case_number }} - {{ case_title }}

**Date:** {{ statement_date }}

---

## STATEMENT OF {{ witness_name }}

I, {{ witness_name }}, make this statement in support of the claim/response in the above matter.

## Background

{{ statement_content }}

## Declaration

I believe that the facts stated in this witness statement are true. I understand that proceedings for contempt of court may be brought against anyone who makes, or causes to be made, a false statement in a document verified by a statement of truth without an honest belief in its truth.

**Signed:** _______________________

**Name:** {{ witness_name }}

**Date:** {{ statement_date }}
"""
    }

    with open(os.path.join(templates_dir, 'witness_statement.json'), 'w') as f:
        json.dump(witness_statement, f, indent=2)

    # Subject Access Request Template
    sar_template = {
        'name': 'Subject Access Request',
        'description': 'GDPR Subject Access Request letter',
        'category': 'Correspondence',
        'fields': [
            {'name': 'recipient_name', 'label': 'Recipient Name', 'type': 'text', 'required': True},
            {'name': 'recipient_address', 'label': 'Recipient Address', 'type': 'textarea', 'required': True},
            {'name': 'data_requested', 'label': 'Data Requested', 'type': 'textarea', 'required': True}
        ],
        'content': """{{ recipient_name }}
{{ recipient_address }}

{{ today }}

Dear Sir/Madam,

## SUBJECT ACCESS REQUEST UNDER GDPR

**Re: {{ case_title }}**

I am writing to make a formal request under the General Data Protection Regulation (GDPR) for access to personal data held about me by your organization.

## Information Requested

{{ data_requested }}

## Legal Basis

This request is made under Article 15 of the GDPR, which gives me the right to obtain:
- Confirmation that my data is being processed
- Access to my personal data
- Other supplementary information

## Time Limit

Please provide the requested information within one month of receipt of this request, as required by the GDPR.

## Format

Please provide the information in electronic format where possible.

Yours faithfully,

{{ primary_contact_name }}
"""
    }

    with open(os.path.join(templates_dir, 'subject_access_request.json'), 'w') as f:
        json.dump(sar_template, f, indent=2)

    # Disclosure Request Template
    disclosure_request = {
        'name': 'Disclosure Request',
        'description': 'Request for disclosure of documents',
        'category': 'Correspondence',
        'fields': [
            {'name': 'recipient_name', 'label': 'Recipient Name', 'type': 'text', 'required': True},
            {'name': 'documents_requested', 'label': 'Documents Requested', 'type': 'textarea', 'required': True},
            {'name': 'deadline_date', 'label': 'Deadline Date', 'type': 'date', 'required': True}
        ],
        'content': """{{ recipient_name }}

{{ today }}

Dear {{ recipient_name }},

## REQUEST FOR DISCLOSURE

**Case:** {{ case_number }} - {{ case_title }}

I write to request disclosure of the following documents:

{{ documents_requested }}

## Deadline

Please provide the requested documents by {{ deadline_date }}.

## Standard Disclosure

This request is made in accordance with the tribunal's standard disclosure obligations. If you believe any documents are not relevant or are privileged, please provide a list of such documents with an explanation.

## Format

Please provide documents in their original format where possible.

I look forward to receiving the requested disclosure by the specified deadline.

Yours sincerely,

{{ primary_contact_name }}
"""
    }

    with open(os.path.join(templates_dir, 'disclosure_request.json'), 'w') as f:
        json.dump(disclosure_request, f, indent=2)

    print("Default templates created successfully")


if __name__ == "__main__":
    create_default_templates()
