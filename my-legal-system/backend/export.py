"""
Export Module - Export case data to Word and PDF formats
"""
from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.lib import colors
from typing import Dict, List, Optional
from datetime import datetime
import os
from sqlalchemy.orm import Session
from models import (
    Case, Milestone, Task, Document, TimelineEvent,
    Contact, Note, Party, CriticalDate
)


class ExportEngine:
    """Export case data to various formats"""

    def __init__(self, db: Session, case_id: int):
        self.db = db
        self.case_id = case_id
        self.case = db.query(Case).filter(Case.id == case_id).first()

        if not self.case:
            raise ValueError(f"Case {case_id} not found")

        # Create exports directory
        self.exports_dir = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),
            'exports'
        )
        os.makedirs(self.exports_dir, exist_ok=True)

    def export_case_summary_word(self) -> Dict:
        """Export complete case summary to Word document"""
        try:
            # Create document
            doc = DocxDocument()

            # Set up styles
            style = doc.styles['Normal']
            font = style.font
            font.name = 'Arial'
            font.size = Pt(11)

            # Title page
            title = doc.add_heading('CASE SUMMARY', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph('')
            doc.add_paragraph(f"Case Number: {self.case.case_number}").bold = True
            doc.add_paragraph(f"Case Title: {self.case.title}").bold = True
            doc.add_paragraph(f"Case Type: {self.case.case_type}")
            doc.add_paragraph(f"Status: {self.case.status}")
            doc.add_paragraph(f"Priority: {self.case.priority}")
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y at %H:%M')}")

            if self.case.mission_statement:
                doc.add_paragraph('')
                p = doc.add_paragraph()
                p.add_run('Mission Statement:').bold = True
                doc.add_paragraph(self.case.mission_statement)

            doc.add_page_break()

            # Milestones section
            self._add_milestones_to_word(doc)

            # Tasks section
            self._add_tasks_to_word(doc)

            # Timeline section
            self._add_timeline_to_word(doc)

            # Documents section
            self._add_documents_to_word(doc)

            # Notes section
            self._add_notes_to_word(doc)

            # Save document
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{self.case.case_number}_summary_{timestamp}.docx"
            file_path = os.path.join(self.exports_dir, filename)

            doc.save(file_path)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'format': 'docx'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def export_case_summary_pdf(self) -> Dict:
        """Export complete case summary to PDF"""
        try:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{self.case.case_number}_summary_{timestamp}.pdf"
            file_path = os.path.join(self.exports_dir, filename)

            # Create PDF
            doc = SimpleDocTemplate(
                file_path,
                pagesize=A4,
                rightMargin=72,
                leftMargin=72,
                topMargin=72,
                bottomMargin=72
            )

            story = []
            styles = getSampleStyleSheet()

            # Add custom styles
            styles.add(ParagraphStyle(
                name='CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#2c3e50'),
                spaceAfter=30,
                alignment=TA_CENTER
            ))

            # Title
            story.append(Paragraph('CASE SUMMARY', styles['CustomTitle']))
            story.append(Spacer(1, 0.3*inch))

            # Case information
            case_info = f"""
            <b>Case Number:</b> {self.case.case_number}<br/>
            <b>Case Title:</b> {self.case.title}<br/>
            <b>Case Type:</b> {self.case.case_type}<br/>
            <b>Status:</b> {self.case.status}<br/>
            <b>Priority:</b> {self.case.priority}<br/>
            <b>Generated:</b> {datetime.now().strftime('%d %B %Y at %H:%M')}
            """
            story.append(Paragraph(case_info, styles['Normal']))

            if self.case.mission_statement:
                story.append(Spacer(1, 0.2*inch))
                story.append(Paragraph('<b>Mission Statement:</b>', styles['Normal']))
                story.append(Paragraph(self.case.mission_statement, styles['Normal']))

            story.append(PageBreak())

            # Add sections
            self._add_milestones_to_pdf(story, styles)
            self._add_tasks_to_pdf(story, styles)
            self._add_timeline_to_pdf(story, styles)
            self._add_documents_to_pdf(story, styles)

            # Build PDF
            doc.build(story)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'format': 'pdf'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def export_timeline_word(self) -> Dict:
        """Export timeline/chronology to Word document"""
        try:
            doc = DocxDocument()

            # Title
            title = doc.add_heading('CASE CHRONOLOGY', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph(f"Case: {self.case.case_number} - {self.case.title}")
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y')}")
            doc.add_paragraph('')

            # Get timeline events
            events = self.db.query(TimelineEvent).filter(
                TimelineEvent.case_id == self.case_id
            ).order_by(TimelineEvent.event_date).all()

            if events:
                # Create table
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Light Grid Accent 1'

                # Header row
                header_cells = table.rows[0].cells
                header_cells[0].text = 'Date'
                header_cells[1].text = 'Event'
                header_cells[2].text = 'Type'
                header_cells[3].text = 'Party'

                # Data rows
                for event in events:
                    row_cells = table.add_row().cells
                    row_cells[0].text = event.event_date.strftime('%d/%m/%Y %H:%M') if event.event_date else ''
                    row_cells[1].text = event.title + ('\n' + event.description if event.description else '')
                    row_cells[2].text = event.event_type or ''
                    row_cells[3].text = event.party_involved or ''
            else:
                doc.add_paragraph('No timeline events recorded.')

            # Save document
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{self.case.case_number}_chronology_{timestamp}.docx"
            file_path = os.path.join(self.exports_dir, filename)

            doc.save(file_path)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'format': 'docx'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def export_notes_word(self) -> Dict:
        """Export case notes to Word document"""
        try:
            doc = DocxDocument()

            # Title
            title = doc.add_heading('CASE NOTES', level=0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph(f"Case: {self.case.case_number} - {self.case.title}")
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y')}")
            doc.add_page_break()

            # Get notes
            notes = self.db.query(Note).filter(
                Note.case_id == self.case_id
            ).order_by(Note.created_date.desc()).all()

            if notes:
                for note in notes:
                    # Note heading
                    doc.add_heading(note.title or note.note_type, level=2)

                    # Metadata
                    meta = doc.add_paragraph()
                    meta.add_run(f"Type: {note.note_type} | ").italic = True
                    meta.add_run(f"Date: {note.created_date.strftime('%d %B %Y')}").italic = True
                    if note.tags:
                        meta.add_run(f" | Tags: {note.tags}").italic = True

                    # Content
                    doc.add_paragraph(note.content)
                    doc.add_paragraph('')  # Spacing
            else:
                doc.add_paragraph('No notes recorded.')

            # Save document
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f"{self.case.case_number}_notes_{timestamp}.docx"
            file_path = os.path.join(self.exports_dir, filename)

            doc.save(file_path)

            return {
                'success': True,
                'file_path': file_path,
                'filename': filename,
                'format': 'docx'
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e)
            }

    def _add_milestones_to_word(self, doc: DocxDocument):
        """Add milestones section to Word document"""
        doc.add_heading('Milestones', level=1)

        milestones = self.db.query(Milestone).filter(
            Milestone.case_id == self.case_id
        ).order_by(Milestone.order_index).all()

        if milestones:
            for milestone in milestones:
                status_icon = '✓' if milestone.is_completed else ('▶' if milestone.completion_percentage > 0 else '□')
                p = doc.add_paragraph(f"{status_icon} ", style='List Bullet')
                p.add_run(f"{milestone.title}").bold = True
                p.add_run(f" ({milestone.completion_percentage}%)")

                details = doc.add_paragraph(style='List Bullet 2')
                details.add_run(f"Type: {milestone.milestone_type} | ")
                if milestone.target_date:
                    details.add_run(f"Target: {milestone.target_date.strftime('%d/%m/%Y')} | ")
                if milestone.is_completed:
                    details.add_run(f"Completed: {milestone.completed_date.strftime('%d/%m/%Y')}")

                if milestone.description:
                    doc.add_paragraph(milestone.description, style='List Bullet 2')
        else:
            doc.add_paragraph('No milestones defined.')

        doc.add_paragraph('')

    def _add_tasks_to_word(self, doc: DocxDocument):
        """Add tasks section to Word document"""
        doc.add_heading('Tasks', level=1)

        tasks = self.db.query(Task).filter(
            Task.case_id == self.case_id
        ).order_by(Task.due_date).all()

        if tasks:
            # Group by status
            pending = [t for t in tasks if t.status != 'Completed']
            completed = [t for t in tasks if t.status == 'Completed']

            if pending:
                doc.add_heading('Pending Tasks', level=2)
                for task in pending:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(f"{task.title} ").bold = True
                    p.add_run(f"[{task.priority}]")

                    if task.due_date:
                        p.add_run(f" - Due: {task.due_date.strftime('%d/%m/%Y')}")

            if completed:
                doc.add_heading('Completed Tasks', level=2)
                for task in completed:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(f"✓ {task.title}")
        else:
            doc.add_paragraph('No tasks defined.')

        doc.add_paragraph('')

    def _add_timeline_to_word(self, doc: DocxDocument):
        """Add timeline section to Word document"""
        doc.add_heading('Timeline', level=1)

        events = self.db.query(TimelineEvent).filter(
            TimelineEvent.case_id == self.case_id
        ).order_by(TimelineEvent.event_date.desc()).limit(20).all()

        if events:
            for event in events:
                key_marker = '⭐ ' if event.is_key_event else ''
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"{key_marker}{event.event_date.strftime('%d/%m/%Y')} - ").bold = True
                p.add_run(event.title)

                if event.description:
                    doc.add_paragraph(event.description, style='List Bullet 2')
        else:
            doc.add_paragraph('No timeline events recorded.')

        doc.add_paragraph('')

    def _add_documents_to_word(self, doc: DocxDocument):
        """Add documents section to Word document"""
        doc.add_heading('Documents', level=1)

        documents = self.db.query(Document).filter(
            Document.case_id == self.case_id
        ).order_by(Document.category, Document.document_date).all()

        if documents:
            # Group by category
            categories = {}
            for doc_item in documents:
                cat = doc_item.category or 'Uncategorized'
                if cat not in categories:
                    categories[cat] = []
                categories[cat].append(doc_item)

            for category, docs in categories.items():
                doc.add_heading(category, level=2)
                for doc_item in docs:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(doc_item.filename).bold = True
                    if doc_item.document_date:
                        p.add_run(f" ({doc_item.document_date.strftime('%d/%m/%Y')})")
        else:
            doc.add_paragraph('No documents uploaded.')

        doc.add_paragraph('')

    def _add_notes_to_word(self, doc: DocxDocument):
        """Add notes section to Word document"""
        doc.add_heading('Notes', level=1)

        notes = self.db.query(Note).filter(
            Note.case_id == self.case_id
        ).order_by(Note.created_date.desc()).limit(10).all()

        if notes:
            for note in notes:
                doc.add_heading(note.title or note.note_type, level=2)
                p = doc.add_paragraph()
                p.add_run(f"{note.created_date.strftime('%d/%m/%Y')} - {note.note_type}").italic = True
                doc.add_paragraph(note.content[:200] + ('...' if len(note.content) > 200 else ''))
        else:
            doc.add_paragraph('No notes recorded.')

    def _add_milestones_to_pdf(self, story: List, styles):
        """Add milestones section to PDF"""
        story.append(Paragraph('MILESTONES', styles['Heading1']))
        story.append(Spacer(1, 0.2*inch))

        milestones = self.db.query(Milestone).filter(
            Milestone.case_id == self.case_id
        ).order_by(Milestone.order_index).all()

        if milestones:
            for milestone in milestones:
                status_icon = '✓' if milestone.is_completed else ('▶' if milestone.completion_percentage > 0 else '□')
                text = f"{status_icon} <b>{milestone.title}</b> ({milestone.completion_percentage}%)"
                story.append(Paragraph(text, styles['Normal']))

                if milestone.description:
                    story.append(Paragraph(milestone.description, styles['Normal']))

                story.append(Spacer(1, 0.1*inch))

        story.append(Spacer(1, 0.3*inch))

    def _add_tasks_to_pdf(self, story: List, styles):
        """Add tasks section to PDF"""
        story.append(Paragraph('TASKS', styles['Heading1']))
        story.append(Spacer(1, 0.2*inch))

        tasks = self.db.query(Task).filter(
            Task.case_id == self.case_id
        ).order_by(Task.due_date).all()

        if tasks:
            for task in tasks:
                status_icon = '✓' if task.status == 'Completed' else '□'
                due_text = f" - Due: {task.due_date.strftime('%d/%m/%Y')}" if task.due_date else ""
                text = f"{status_icon} <b>{task.title}</b> [{task.priority}]{due_text}"
                story.append(Paragraph(text, styles['Normal']))
                story.append(Spacer(1, 0.05*inch))

        story.append(Spacer(1, 0.3*inch))

    def _add_timeline_to_pdf(self, story: List, styles):
        """Add timeline section to PDF"""
        story.append(Paragraph('TIMELINE', styles['Heading1']))
        story.append(Spacer(1, 0.2*inch))

        events = self.db.query(TimelineEvent).filter(
            TimelineEvent.case_id == self.case_id
        ).order_by(TimelineEvent.event_date.desc()).limit(20).all()

        if events:
            for event in events:
                key_marker = '⭐ ' if event.is_key_event else ''
                text = f"{key_marker}<b>{event.event_date.strftime('%d/%m/%Y')}</b> - {event.title}"
                story.append(Paragraph(text, styles['Normal']))
                if event.description:
                    story.append(Paragraph(event.description, styles['Normal']))
                story.append(Spacer(1, 0.05*inch))

        story.append(Spacer(1, 0.3*inch))

    def _add_documents_to_pdf(self, story: List, styles):
        """Add documents section to PDF"""
        story.append(Paragraph('DOCUMENTS', styles['Heading1']))
        story.append(Spacer(1, 0.2*inch))

        documents = self.db.query(Document).filter(
            Document.case_id == self.case_id
        ).order_by(Document.category).all()

        if documents:
            # Group by category
            categories = {}
            for doc in documents:
                cat = doc.category or 'Uncategorized'
                if cat not in categories:
                    categories[cat] = []
                categories[cat].append(doc)

            for category, docs in categories.items():
                story.append(Paragraph(f'<b>{category}</b>', styles['Heading2']))
                for doc in docs:
                    date_text = f" ({doc.document_date.strftime('%d/%m/%Y')})" if doc.document_date else ""
                    story.append(Paragraph(f"• {doc.filename}{date_text}", styles['Normal']))
                story.append(Spacer(1, 0.1*inch))
